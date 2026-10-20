"""TC-UNIT-07 - exporters: tex / pdf / docx / html / md."""
from __future__ import annotations

import pytest

from pagex.core.blocks import Block, block_id_for, make_document, make_page
from pagex.core import latex_export as lx


def make_doc(kind_note: bool = True) -> "object":
    doc = make_document("doc_expo", "Export Demo & Review", kind="image")
    page = make_page(doc, 1, width=1240, height=1754, image_path="p.png")
    page.caption = "A synthetic page used by the exporter tests."
    page.blocks = [
        Block(block_id=block_id_for(doc.doc_id, 0), page_id=page.page_id,
              doc_id=doc.doc_id, seq=0, kind="heading", page_no=1,
              bbox=(70, 40, 1170, 120), text="Results & Discussion"),
        Block(block_id=block_id_for(doc.doc_id, 1), page_id=page.page_id,
              doc_id=doc.doc_id, seq=1, kind="paragraph", page_no=1,
              bbox=(70, 140, 1170, 320),
              text="Energy rises with mass (50% & 75% loads)."),
        Block(block_id=block_id_for(doc.doc_id, 2), page_id=page.page_id,
              doc_id=doc.doc_id, seq=2, kind="math", page_no=1,
              bbox=(80, 340, 700, 440), latex=r"E = \frac{1}{2} m v^2"),
        Block(block_id=block_id_for(doc.doc_id, 3), page_id=page.page_id,
              doc_id=doc.doc_id, seq=3, kind="table", page_no=1,
              bbox=(80, 480, 1100, 700),
              latex="\\begin{tabular}{lll}\nTrial & Mass & Energy \\\\\n"
                    "1 & 0.50 & 0.31 \\\\\n2 & 0.75 & 0.49 \\\\\n\\end{tabular}"),
        Block(block_id=block_id_for(doc.doc_id, 4), page_id=page.page_id,
              doc_id=doc.doc_id, seq=4, kind="diagram", page_no=1,
              bbox=(80, 740, 900, 1000), latex=r"\draw (0,0) -- (2,0);"),
        Block(block_id=block_id_for(doc.doc_id, 5), page_id=page.page_id,
              doc_id=doc.doc_id, seq=5, kind="caption", page_no=1,
              bbox=(70, 1040, 1170, 1100),
              text="Table 1: measured energy."),
    ]
    return doc


def test_clean_tex_has_no_markers_but_keeps_content():
    doc = make_doc()
    clean = lx.document_to_tex(doc, annotated=False)
    assert "%PX:" not in clean
    # '&' must be escaped or TeX would treat it as a column separator
    assert "\\section*{Results \\& Discussion}" in clean
    assert clean.startswith("\\documentclass")


def test_annotated_tex_keeps_markers():
    doc = make_doc()
    tex = lx.document_to_tex(doc, annotated=True)
    assert "%PX:block=doc_expo-b0000;page=1;kind=heading" in tex
    assert "\\section*{Results \\& Discussion}" in tex


def test_reportlab_pdf_writes_magic_bytes(tmp_path):
    """No TeX installed -> the fallback must still produce a real PDF."""
    doc = make_doc()
    out = tmp_path / "out.pdf"
    lx.reportlab_pdf(doc, out)
    assert out.exists() and out.stat().st_size > 500
    assert out.read_bytes()[:4] == b"%PDF"


def test_export_pdf_uses_fallback_when_pdflatex_missing(tmp_path, monkeypatch):
    monkeypatch.setattr(lx, "has_pdflatex", lambda: False)
    doc = make_doc()
    out = lx.export_document(doc, tmp_path / "doc.pdf", "pdf")
    assert out.read_bytes()[:4] == b"%PDF"


def test_faithful_pdf_keeps_page_geometry_and_text_layer(tmp_path):
    """The same-layout export keeps the page size and stays searchable."""
    fitz = pytest.importorskip("fitz")
    doc = make_doc()
    out = tmp_path / "faithful.pdf"
    lx.export_document(doc, out, "pdf", pdf_layout="faithful")
    assert out.read_bytes()[:4] == b"%PDF"

    with fitz.open(out) as pdf:
        assert pdf.page_count == 1
        rect = pdf[0].rect
        # 1240 x 1754 px at 200 dpi, expressed in PDF points
        assert rect.width == pytest.approx(1240 * 72 / 200, rel=1e-3)
        assert rect.height == pytest.approx(1754 * 72 / 200, rel=1e-3)
        text = pdf[0].get_text()
    assert "Results & Discussion" in text
    assert "Energy rises with mass" in text
    assert "Trial" in text


def test_export_document_rejects_unknown_pdf_layout(tmp_path):
    with pytest.raises(ValueError):
        lx.export_document(make_doc(), tmp_path / "x.pdf", "pdf",
                           pdf_layout="scrambled")


def test_docx_roundtrip_read_back(tmp_path):
    from docx import Document as DocxDocument
    doc = make_doc()
    out = lx.export_document(doc, tmp_path / "doc.docx", "docx")
    assert out.exists()
    read = DocxDocument(str(out))
    texts = [p.text for p in read.paragraphs]
    assert any("Results & Discussion" in t for t in texts)
    assert any("Energy rises with mass" in t for t in texts)
    # the table made it across
    assert len(read.tables) >= 1
    cells = [c.text for row in read.tables[0].rows for c in row.cells]
    assert "Trial" in cells and "0.31" in cells


def test_html_preview_contains_blocks(tmp_path):
    doc = make_doc()
    html = lx.document_to_html(doc, show_markers=True)
    assert "<h2>Results &amp; Discussion</h2>" in html
    assert "E = \\frac{1}{2} m v^2" in html or "frac" in html
    assert "<table>" in html
    assert "%PX:block=" in html
    assert html.startswith("<!doctype html>")


def test_markdown_export():
    doc = make_doc()
    md = lx.document_to_markdown(doc)
    assert md.startswith("# Export Demo")
    assert "## Results & Discussion" in md
    assert "$$" in md
    assert "| Trial |" in md


def test_export_json_is_parseable(tmp_path):
    import json
    doc = make_doc()
    out = lx.export_document(doc, tmp_path / "doc.json", "json")
    data = json.loads(out.read_text())
    assert data["pages"][0]["blocks"][0]["kind"] == "heading"
    assert data["kinds"]["math"] == 1


def test_export_unknown_format_raises(tmp_path):
    with pytest.raises(ValueError):
        lx.export_document(make_doc(), tmp_path / "x.thing", "thing")


def test_tabular_to_md_handles_single_column():
    md = lx._tabular_to_md("\\begin{tabular}{l}\nonly \\\\\n\\end{tabular}")
    assert md.splitlines()[0].startswith("| only")


def test_escape_xml():
    assert lx.escape_xml("a<b&c>") == "a&lt;b&amp;c&gt;"
